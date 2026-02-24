"""
خدمات تصدير قوائم الأسعار
Price List Export Services

يشمل:
- تصدير PDF بتصميم احترافي
- تصدير Excel
- تصدير CSV
"""

import os
import io
from decimal import Decimal
from typing import Dict, List, Optional
from datetime import datetime
from django.conf import settings
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.utils import timezone


class PriceListExporter:
    """خدمة تصدير قوائم الأسعار"""
    
    def __init__(self, price_list, families=None, include_tax_column=True, 
                 show_cost=False, show_margin=False):
        """
        Args:
            price_list: PriceList instance
            families: قائمة العائلات (فارغ = الكل)
            include_tax_column: إظهار عمود السعر شامل الضريبة
            show_cost: إظهار التكلفة (للاستخدام الداخلي)
            show_margin: إظهار هامش الربح (للاستخدام الداخلي)
        """
        self.price_list = price_list
        self.families = families
        self.include_tax_column = include_tax_column
        self.show_cost = show_cost
        self.show_margin = show_margin
        
        self._load_data()
    
    def _load_data(self):
        """تحميل البيانات"""
        from .models_price_list import (
            ProductFamily, ProductSize, ProductVariantPrice, 
            PriceListSettings
        )
        from .services_price_list import PriceListService
        
        # تحميل الإعدادات
        self.settings = PriceListSettings.get_settings()
        
        # تحميل العائلات
        if self.families is None:
            self.families = ProductFamily.objects.filter(
                is_active=True,
                show_in_price_list=True
            ).order_by('sort_order', 'name')
        
        # تحميل المقاسات
        self.sizes = ProductSize.objects.filter(
            is_active=True,
            product_families__in=self.families
        ).distinct().order_by('sort_order', 'width', 'length')
        
        # استخراج العروض الفريدة فقط (الطول لا يؤثر على السعر)
        self.unique_widths = sorted(set(size.width for size in self.sizes))
        
        # بناء مصفوفة الأسعار
        self.price_matrix = PriceListService.generate_price_matrix(
            self.price_list, 
            self.families
        )
        
        # بناء قاموس للبحث السريع عن الأسعار بالعرض فقط
        self._build_width_price_map()
    
    def _build_width_price_map(self):
        """
        بناء قاموس للبحث السريع عن الأسعار بالعرض فقط
        لأن الطول لا يؤثر على السعر
        """
        self.width_price_map = {}
        
        for family_data in self.price_matrix:
            family_id = family_data['family_id']
            self.width_price_map[family_id] = {}
            
            for size_key, price_info in family_data['prices'].items():
                # استخراج العرض من size_key (مثال: "90x190" -> 90)
                width = int(size_key.split('x')[0])
                # أخذ أول سعر موجود لهذا العرض (لأن كل الأطوال لها نفس السعر)
                if width not in self.width_price_map[family_id]:
                    self.width_price_map[family_id][width] = price_info
    
    def _get_price_by_width(self, family_id, width):
        """الحصول على السعر بناءً على العرض فقط"""
        if family_id in self.width_price_map:
            return self.width_price_map[family_id].get(width)
        return None
    
    def export_to_excel(self) -> HttpResponse:
        """تصدير إلى Excel - تصميم مطابق لـ Rita"""
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ImportError("openpyxl مطلوب للتصدير. قم بتثبيته: pip install openpyxl")
        
        # الحصول على بيانات الشركة
        company_data = self.settings.get_all_company_data()
        company_name = company_data['name']
        company_name_en = company_data.get('name_en', '')
        company_phone = company_data.get('phone', '')
        company_email = company_data.get('email', '')
        company_website = company_data.get('website', '')
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "قائمة الأسعار"
        ws.sheet_view.rightToLeft = True  # دعم RTL
        
        # الألوان
        header_fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        subheader_fill = PatternFill(start_color="2980B9", end_color="2980B9", fill_type="solid")
        subheader_font = Font(bold=True, color="FFFFFF", size=9)
        alternate_fill = PatternFill(start_color="ECF0F1", end_color="ECF0F1", fill_type="solid")
        size_col_fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        num_products = len(list(self.families))
        total_cols = num_products + 2  # المقاس + المنتجات + الأوصاف
        
        # الترويسة - اسم الشركة
        row = 1
        ws.merge_cells(f'A{row}:' + get_column_letter(total_cols) + str(row))
        ws[f'A{row}'] = f"{company_name}  |  {company_name_en or company_name}"
        ws[f'A{row}'].font = Font(bold=True, size=20, color="1A5276")
        ws[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[row].height = 35
        
        # تاريخ قائمة الأسعار
        row += 1
        ws.merge_cells(f'A{row}:' + get_column_letter(total_cols) + str(row))
        date_str = timezone.now().strftime('%d/%m/%Y')
        ws[f'A{row}'] = f"PRICE LIST STARTING FROM {date_str} - قائمة الأسعار: {self.price_list.name}"
        ws[f'A{row}'].font = Font(bold=True, size=11, color="FFFFFF")
        ws[f'A{row}'].fill = PatternFill(start_color="F39C12", end_color="F39C12", fill_type="solid")
        ws[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[row].height = 25
        
        row += 1  # صف فارغ
        
        # رؤوس الأعمدة - الصف الأول (أسماء المنتجات)
        row += 1
        header_row = row
        
        # عمود المقاس
        ws[f'A{row}'] = "المقاس"
        ws[f'A{row}'].fill = header_fill
        ws[f'A{row}'].font = header_font
        ws[f'A{row}'].border = border
        ws[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
        
        # أسماء المنتجات
        col = 2
        families_list = list(self.families)
        for family in families_list:
            cell = ws.cell(row=row, column=col)
            cell.value = family.name
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            col += 1
        
        # عمود الأوصاف
        cell = ws.cell(row=row, column=col)
        cell.value = "الأوصاف"
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # الصف الثاني - الارتفاعات
        row += 1
        ws[f'A{row}'] = "العرض"
        ws[f'A{row}'].fill = subheader_fill
        ws[f'A{row}'].font = subheader_font
        ws[f'A{row}'].border = border
        ws[f'A{row}'].alignment = Alignment(horizontal='center', vertical='center')
        
        col = 2
        for family in families_list:
            cell = ws.cell(row=row, column=col)
            height = getattr(family, 'default_height', 30) or 30
            cell.value = f"{height} cm"
            cell.fill = subheader_fill
            cell.font = subheader_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            col += 1
        
        cell = ws.cell(row=row, column=col)
        cell.value = "المقاس"
        cell.fill = subheader_fill
        cell.font = subheader_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # تعيين عرض الأعمدة
        ws.column_dimensions['A'].width = 12
        for i in range(2, total_cols + 1):
            ws.column_dimensions[get_column_letter(i)].width = 10
        ws.column_dimensions[get_column_letter(total_cols)].width = 15
        
        row += 1
        
        # البيانات (العروض الفريدة كصفوف - الطول لا يؤثر على السعر)
        for idx, width in enumerate(self.unique_widths):
            # عمود المقاس (العرض)
            cell_a = ws[f'A{row}']
            cell_a.value = width
            cell_a.fill = size_col_fill
            cell_a.font = Font(bold=True, color="FFFFFF")
            cell_a.border = border
            cell_a.alignment = Alignment(horizontal='center', vertical='center')
            
            # الأسعار لكل منتج
            col = 2
            for family_data in self.price_matrix:
                cell = ws.cell(row=row, column=col)
                
                price_info = self._get_price_by_width(family_data['family_id'], width)
                if price_info:
                    price = price_info['price_with_tax'] if self.price_list.includes_tax else price_info['price']
                    cell.value = float(price)
                    cell.number_format = '#,##0'
                else:
                    cell.value = '-'
                
                cell.border = border
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.font = Font(bold=True, size=10)
                if idx % 2 == 1:
                    cell.fill = alternate_fill
                
                col += 1
            
            # عمود الأوصاف (الأطوال المتاحة)
            cell = ws.cell(row=row, column=col)
            cell.value = "190/195/200"
            cell.fill = size_col_fill
            cell.font = Font(bold=True, color="FFFFFF", size=9)
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            
            row += 1
        
        # التذييل
        row += 1
        ws.merge_cells(f'A{row}:' + get_column_letter(total_cols) + str(row))
        contact_info = []
        if company_phone:
            contact_info.append(f"📞 {company_phone}")
        if company_website:
            contact_info.append(f"🌐 {company_website}")
        if company_email:
            contact_info.append(f"📧 {company_email}")
        ws[f'A{row}'] = " | ".join(contact_info)
        ws[f'A{row}'].alignment = Alignment(horizontal='center')
        ws[f'A{row}'].font = Font(size=10)
        
        # إنشاء الاستجابة
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"price_list_{self.price_list.code}_{timezone.now().strftime('%Y%m%d')}.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        return response
    
    def export_to_pdf(self) -> HttpResponse:
        """تصدير إلى PDF"""
        try:
            from weasyprint import HTML, CSS
        except ImportError:
            # استخدام xhtml2pdf كبديل
            return self._export_to_pdf_xhtml2pdf()
        
        html_content = self._generate_html()
        
        # تحويل HTML إلى PDF
        html = HTML(string=html_content, base_url=settings.BASE_DIR)
        
        # CSS إضافي للطباعة
        css = CSS(string='''
            @page {
                size: A4 landscape;
                margin: 1cm;
            }
            body {
                direction: rtl;
                font-family: 'Cairo', 'Noto Sans Arabic', Arial, sans-serif;
            }
        ''')
        
        pdf = html.write_pdf(stylesheets=[css])
        
        response = HttpResponse(pdf, content_type='application/pdf')
        filename = f"price_list_{self.price_list.code}_{timezone.now().strftime('%Y%m%d')}.pdf"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        return response
    
    def _export_to_pdf_xhtml2pdf(self) -> HttpResponse:
        """تصدير PDF باستخدام xhtml2pdf"""
        try:
            from xhtml2pdf import pisa
            from xhtml2pdf.default import DEFAULT_FONT
        except ImportError:
            raise ImportError("xhtml2pdf أو weasyprint مطلوب للتصدير. قم بتثبيت أحدهما")
        
        # تسجيل الخطوط العربية
        try:
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            
            # تسجيل خط Noto Naskh Arabic
            pdfmetrics.registerFont(TTFont('NotoArabic', '/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf'))
            pdfmetrics.registerFont(TTFont('NotoArabicBold', '/usr/share/fonts/truetype/noto/NotoNaskhArabic-Bold.ttf'))
            
            # تسجيل خط DejaVu Sans كـ fallback
            pdfmetrics.registerFont(TTFont('DejaVuSans', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
        except Exception as e:
            print(f"Warning: Could not register Arabic fonts: {e}")
        
        html_content = self._generate_html()
        
        # إصلاح النص العربي للـ PDF
        html_content = self._fix_arabic_text(html_content)
        
        response = HttpResponse(content_type='application/pdf')
        filename = f"price_list_{self.price_list.code}_{timezone.now().strftime('%Y%m%d')}.pdf"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        pisa_status = pisa.CreatePDF(html_content, dest=response)
        
        if pisa_status.err:
            return HttpResponse('خطأ في إنشاء PDF', status=500)
        
        return response
    
    def _fix_arabic_text(self, html_content: str) -> str:
        """إصلاح النص العربي ليظهر بشكل صحيح في PDF"""
        try:
            import arabic_reshaper
            from bidi.algorithm import get_display
            import re
            
            # البحث عن النصوص العربية وإصلاحها
            def reshape_arabic(match):
                text = match.group(0)
                # تجاهل الأرقام والتواريخ
                if re.match(r'^[\d\s/\-\.:]+$', text):
                    return text
                # إعادة تشكيل النص العربي
                reshaped = arabic_reshaper.reshape(text)
                bidi_text = get_display(reshaped)
                return bidi_text
            
            # نمط للبحث عن النص العربي
            arabic_pattern = r'[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]+'
            
            # استبدال كل النصوص العربية
            fixed_html = re.sub(arabic_pattern, reshape_arabic, html_content)
            
            return fixed_html
        except Exception as e:
            print(f"Warning: Could not fix Arabic text: {e}")
            return html_content
    
    def _generate_html(self) -> str:
        """توليد HTML لقائمة الأسعار - تصميم مطابق لـ Rita"""
        
        # الحصول على بيانات الشركة
        company_data = self.settings.get_all_company_data()
        company_name = company_data['name']
        company_name_en = company_data.get('name_en', '')
        company_logo = company_data.get('logo')
        company_phone = company_data.get('phone', '')
        company_email = company_data.get('email', '')
        company_website = company_data.get('website', '')
        
        # تحديد مسار اللوجو
        logo_url = ''
        if company_logo:
            try:
                logo_url = company_logo.url
            except:
                logo_url = ''
        
        # تاريخ قائمة الأسعار
        date_str = timezone.now().strftime('%d/%m/%Y')
        
        html = f'''
<!DOCTYPE html>
<html dir="rtl" lang="ar">
<head>
    <meta charset="UTF-8">
    <style>
        @font-face {{
            font-family: 'NotoArabic';
            src: url('/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf');
        }}
        
        @page {{
            size: A4 landscape;
            margin: 0.4cm;
        }}
        
        body {{
            font-family: 'NotoArabic', 'Noto Naskh Arabic', 'DejaVu Sans', Arial, sans-serif;
            direction: rtl;
            margin: 0;
            padding: 5px;
            background-color: white;
            color: #333;
            font-size: 9px;
        }}
        
        .container {{
            width: 100%;
        }}
        
        /* Header with logos on sides */
        .header-table {{
            width: 100%;
            border-collapse: collapse;
            margin-bottom: 3px;
        }}
        
        .header-table td {{
            vertical-align: middle;
            padding: 3px;
        }}
        
        .logo-cell {{
            width: 50px;
            text-align: center;
        }}
        
        .logo-cell img {{
            width: 45px;
            height: auto;
        }}
        
        .title-cell {{
            text-align: center;
            background: linear-gradient(135deg, #1a5276, #2980b9);
            background-color: #1a5276;
            color: white;
            padding: 8px;
        }}
        
        .company-name {{
            font-size: 18px;
            font-weight: bold;
            margin-bottom: 2px;
        }}
        
        /* Orange date bar */
        .date-bar {{
            background-color: #f39c12;
            color: white;
            text-align: center;
            padding: 3px;
            font-weight: bold;
            font-size: 9px;
            margin-bottom: 3px;
        }}
        
        /* Price table */
        .price-table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 8px;
        }}
        
        .price-table th {{
            background-color: #eef4f9;
            color: #1a5276;
            padding: 4px 2px;
            font-weight: bold;
            font-size: 8px;
            border: 1px solid #bdc3c7;
            text-align: center;
        }}
        
        .price-table td {{
            padding: 3px 2px;
            border: 1px solid #bdc3c7;
            text-align: center;
            font-weight: bold;
            font-size: 8px;
        }}
        
        .price-table tr:nth-child(even) td {{
            background-color: #f9f9f9;
        }}
        
        .size-cell {{
            background-color: #eef4f9;
            color: #1a5276;
            font-weight: bold;
            font-size: 8px;
        }}
        
        .height-row th {{
            background-color: #f39c12;
            color: white;
            font-size: 7px;
        }}
        
        /* Footer */
        .footer {{
            margin-top: 5px;
            padding: 4px;
            background-color: #1a5276;
            color: white;
            font-size: 7px;
        }}
        
        .footer-table {{
            width: 100%;
            border-collapse: collapse;
        }}
        
        .footer-table td {{
            padding: 2px 8px;
            vertical-align: middle;
        }}
        
        .footer-left {{
            text-align: right;
        }}
        
        .footer-center {{
            text-align: center;
        }}
        
        .footer-right {{
            text-align: left;
        }}
    </style>
</head>
<body>
    <div class="container">
        <!-- Header with logos -->
        <table class="header-table">
            <tr>
                <td class="logo-cell">
                    {'<img src="' + logo_url + '" />' if logo_url else ''}
                </td>
                <td class="title-cell">
                    <div class="company-name">{company_name} {company_name_en}</div>
                </td>
                <td class="logo-cell">
                    {'<img src="' + logo_url + '" />' if logo_url else ''}
                </td>
            </tr>
        </table>
        
        <!-- Date bar -->
        <div class="date-bar">
            PRICE LIST STARTING FROM {date_str}
        </div>
        
        <table class="price-table">
            <thead>
                <tr>
'''
        
        # إضافة رؤوس المنتجات
        families_list = list(self.families)
        for family in families_list:
            html += f'<th>{family.name}</th>'
        
        # عمود الأوصاف (الصنف) - على اليمين
        html += '<th>الصنف</th>'
        
        html += '''
                </tr>
                <tr class="height-row">
'''
        # صف الارتفاعات مع "سم"
        for family in families_list:
            height = getattr(family, 'default_height', 30) or 30
            html += f'<th>{height} سم</th>'
        html += '<th>المقاس</th>'
        
        html += '''
                </tr>
            </thead>
            <tbody>
'''
        
        # بناء الصفوف (العروض الفريدة فقط - الطول لا يؤثر على السعر)
        for width in self.unique_widths:
            html += f'''
                <tr>
'''
            
            # إضافة الأسعار لكل منتج
            for family_data in self.price_matrix:
                price_info = self._get_price_by_width(family_data['family_id'], width)
                if price_info:
                    price = price_info['price_with_tax'] if self.price_list.includes_tax else price_info['price']
                    html += f'<td>{int(price)}</td>'
                else:
                    html += '<td>-</td>'
            
            # عمود الأوصاف (العرض + الأطوال المتاحة مع فواصل)
            html += f'<td class="size-cell">{width}<br><span style="font-size: 6px;">190/195/200</span></td>'
            
            html += '''
                </tr>
'''
        
        html += f'''
            </tbody>
        </table>
        
        <!-- Footer like Rita -->
        <div class="footer">
            <table class="footer-table">
                <tr>
                    <td class="footer-right" style="width: 30%;">
                        {'غير شاملة ضريبة القيمة المضافة ' + str(self.price_list.tax_rate) + '%' if not self.price_list.includes_tax else 'شاملة ضريبة القيمة المضافة ' + str(self.price_list.tax_rate) + '%'}
                    </td>
                    <td class="footer-center" style="width: 50%;">
                        {'تليفون: ' + company_phone + ' | ' if company_phone else ''}
                        {company_website + ' | ' if company_website else ''}
                        {company_email if company_email else ''}
                        {' | واتساب: ' + company_data.get('whatsapp', '') if company_data.get('whatsapp') else ''}
                    </td>
                    <td class="footer-left" style="width: 20%;">
                        {company_data.get('address', '') if company_data.get('address') else 'ادارة المبيعات'}
                    </td>
                </tr>
            </table>
        </div>
    </div>
</body>
</html>
'''
        
        return html
    
    def export_to_csv(self) -> HttpResponse:
        """تصدير إلى CSV"""
        import csv
        
        response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
        filename = f"price_list_{self.price_list.code}_{timezone.now().strftime('%Y%m%d')}.csv"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        # إضافة BOM للـ UTF-8 (لدعم Excel)
        response.write('\ufeff')
        
        writer = csv.writer(response)
        
        # الترويسة
        writer.writerow([self.settings.get_company_name()])
        writer.writerow([f"قائمة الأسعار: {self.price_list.name}"])
        writer.writerow([f"تاريخ الإصدار: {timezone.now().strftime('%Y/%m/%d')}"])
        writer.writerow([])
        
        # رؤوس الأعمدة (المنتجات + عمود المقاس)
        families_list = list(self.families)
        headers = [f.name for f in families_list] + ['المقاس', 'الأوصاف']
        writer.writerow(headers)
        
        # صف الارتفاعات
        heights = [f"{getattr(f, 'default_height', 30) or 30} cm" for f in families_list]
        writer.writerow(heights + ['العرض', 'الطول'])
        
        # البيانات (العروض الفريدة)
        for width in self.unique_widths:
            row = []
            
            for family_data in self.price_matrix:
                price_info = self._get_price_by_width(family_data['family_id'], width)
                if price_info:
                    price = price_info['price_with_tax'] if self.price_list.includes_tax else price_info['price']
                    row.append(int(price))
                else:
                    row.append('-')
            
            row.append(width)  # المقاس (العرض)
            row.append('190/195/200')  # الأطوال المتاحة
            
            writer.writerow(row)
        
        return response
    
    def save_export(self, format: str = 'pdf') -> str:
        """حفظ التصدير وتسجيله"""
        from .models_price_list import PriceListExport
        from django.core.files.base import ContentFile
        
        if format == 'excel':
            response = self.export_to_excel()
            ext = 'xlsx'
        elif format == 'csv':
            response = self.export_to_csv()
            ext = 'csv'
        else:
            response = self.export_to_pdf()
            ext = 'pdf'
        
        # إنشاء سجل التصدير
        export_record = PriceListExport.objects.create(
            price_list=self.price_list,
            format=format,
            include_cost=self.show_cost,
            include_margin=self.show_margin,
            include_tax_breakdown=self.include_tax_column,
        )
        
        # حفظ العائلات المصدرة
        if self.families:
            export_record.families.set(self.families)
        
        # حفظ الملف
        filename = f"price_list_{self.price_list.code}_{timezone.now().strftime('%Y%m%d%H%M%S')}.{ext}"
        export_record.file.save(filename, ContentFile(response.content))
        
        return export_record.file.url


def export_all_price_lists(format: str = 'excel') -> HttpResponse:
    """تصدير جميع قوائم الأسعار في ملف واحد"""
    from .models_price_list import PriceList
    
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    except ImportError:
        raise ImportError("openpyxl مطلوب. قم بتثبيته: pip install openpyxl")
    
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # إزالة الورقة الافتراضية
    
    price_lists = PriceList.objects.filter(is_active=True).order_by('name')
    
    for price_list in price_lists:
        exporter = PriceListExporter(price_list)
        
        ws = wb.create_sheet(title=price_list.name[:31])  # حد الـ 31 حرف
        ws.sheet_view.rightToLeft = True
        
        # نسخ البيانات من exporter
        # (تبسيط - يمكن تحسينه لنسخ التنسيق كاملاً)
        row = 1
        ws[f'A{row}'] = f"قائمة الأسعار: {price_list.name}"
        ws[f'A{row}'].font = Font(bold=True, size=14)
        
        row += 2
        
        # الرؤوس
        ws[f'A{row}'] = "المنتج"
        col = 2
        for size in exporter.sizes:
            ws.cell(row=row, column=col, value=str(size))
            col += 1
        
        row += 1
        
        # البيانات
        for family_data in exporter.price_matrix:
            ws[f'A{row}'] = family_data['family_name']
            col = 2
            for size in exporter.sizes:
                size_key = f"{size.width}x{size.length}"
                if size_key in family_data['prices']:
                    price = family_data['prices'][size_key]['price']
                    ws.cell(row=row, column=col, value=float(price))
                col += 1
            row += 1
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"all_price_lists_{timezone.now().strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    return response
