"""
خدمة التصدير الشاملة — RITA ERP Sprint 25
Excel + PDF لكل الجداول والتقارير
"""
import io
from datetime import datetime
from django.http import HttpResponse
from django.template.loader import render_to_string


# ─────────────────────────────────────────────────────────────
# Excel Export
# ─────────────────────────────────────────────────────────────

def export_to_excel(queryset, columns: list, filename: str = 'export', sheet_name: str = 'البيانات') -> HttpResponse:
    """
    تصدير أي queryset إلى Excel.
    
    :param queryset: Django queryset أو list of dicts
    :param columns: قائمة دواميس {'header': 'اسم العمود', 'field': 'اسم الحقل أو callable'}
    :param filename: اسم الملف (بدون امتداد)
    :param sheet_name: اسم الورقة
    """
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.sheet_view.rightToLeft = True  # RTL للعربي

    # ── Header Style ──
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='1F3864', end_color='1F3864', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    # ── Title Row ──
    ws.row_dimensions[1].height = 35
    ws.merge_cells(f'A1:{get_column_letter(len(columns))}1')
    title_cell = ws['A1']
    title_cell.value = f'{sheet_name} — {datetime.now().strftime("%Y-%m-%d %H:%M")}'
    title_cell.font = Font(bold=True, size=14, color='FFFFFF')
    title_cell.fill = PatternFill(start_color='2E4057', end_color='2E4057', fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')

    # ── Column Headers ──
    header_row = 2
    for col_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=col['header'])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = col.get('width', 20)
    ws.row_dimensions[header_row].height = 25

    # ── Data Rows ──
    alt_fill = PatternFill(start_color='EBF3FB', end_color='EBF3FB', fill_type='solid')
    data_align = Alignment(horizontal='right', vertical='center', wrap_text=False)

    for row_idx, obj in enumerate(queryset, start=header_row + 1):
        for col_idx, col in enumerate(columns, start=1):
            field = col['field']
            if callable(field):
                value = field(obj)
            elif isinstance(obj, dict):
                value = obj.get(field, '')
            else:
                value = _get_nested_attr(obj, field)

            # تحويل القيم
            if hasattr(value, 'strftime'):
                value = value.strftime('%Y-%m-%d %H:%M') if hasattr(value, 'hour') else value.strftime('%Y-%m-%d')

            cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if value is not None else '')
            cell.alignment = data_align
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill
        ws.row_dimensions[row_idx].height = 20

    # ── Freeze Header ──
    ws.freeze_panes = f'A{header_row + 1}'

    # ── Auto-filter ──
    ws.auto_filter.ref = f'A{header_row}:{get_column_letter(len(columns))}{header_row}'

    # ── حفظ ──
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx"'
    return response


def _get_nested_attr(obj, field: str):
    """الوصول للحقول المتداخلة مثل: 'customer.name'"""
    parts = field.split('.')
    value = obj
    for part in parts:
        if value is None:
            return ''
        if isinstance(value, dict):
            value = value.get(part, '')
        else:
            value = getattr(value, part, '')
            if callable(value):
                value = value()
    return value if value is not None else ''


# ─────────────────────────────────────────────────────────────
# PDF Export
# ─────────────────────────────────────────────────────────────

def export_to_pdf(template_name: str, context: dict, filename: str = 'export') -> HttpResponse:
    """
    تصدير أي template إلى PDF باستخدام xhtml2pdf.
    
    :param template_name: مسار الـ template
    :param context: بيانات الـ template
    :param filename: اسم الملف (بدون امتداد)
    """
    from xhtml2pdf import pisa

    html_string = render_to_string(template_name, context)
    output = io.BytesIO()
    pisa_status = pisa.CreatePDF(html_string, dest=output, encoding='UTF-8')

    if pisa_status.err:
        raise ValueError(f'خطأ في توليد PDF: {pisa_status.err}')

    output.seek(0)
    response = HttpResponse(output.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = (
        f'attachment; filename="{filename}_{datetime.now().strftime("%Y%m%d_%H%M")}.pdf"'
    )
    return response


def export_queryset_to_pdf(queryset, columns: list, title: str, filename: str = 'export') -> HttpResponse:
    """
    تصدير قائمة بيانات مباشرة إلى PDF (بدون template مخصص).
    """
    rows = []
    for obj in queryset:
        row = []
        for col in columns:
            field = col['field']
            if callable(field):
                value = field(obj)
            elif isinstance(obj, dict):
                value = obj.get(field, '')
            else:
                value = _get_nested_attr(obj, field)
            row.append(str(value) if value is not None else '')
        rows.append(row)

    context = {
        'title': title,
        'headers': [col['header'] for col in columns],
        'rows': rows,
        'generated_at': datetime.now().strftime('%Y-%m-%d %H:%M'),
    }
    return export_to_pdf('exports/generic_table_pdf.html', context, filename)
