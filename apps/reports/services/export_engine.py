"""
محرك التصدير — PDF + Excel
"""
import io
from decimal import Decimal
from django.http import HttpResponse
from django.template.loader import render_to_string
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """تصدير Excel"""

    @classmethod
    def export(cls, title, headers, data, filename='report.xlsx'):
        """
        تصدير بيانات إلى Excel

        headers = ['الكود', 'الاسم', 'الكمية', 'السعر']
        data = [
            ['P001', 'منتج 1', 100, 500],
            ['P002', 'منتج 2', 50, 300],
        ]
        """
        wb = Workbook()
        ws = wb.active
        ws.title = title[:31]  # Excel limit
        ws.sheet_view.rightToLeft = True  # RTL

        # العنوان
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        title_cell = ws.cell(row=1, column=1, value=title)
        title_cell.font = Font(name='Arial', size=16, bold=True)
        title_cell.alignment = Alignment(horizontal='center')

        # التاريخ
        from django.utils import timezone
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        date_cell = ws.cell(row=2, column=1, value=f'تاريخ التصدير: {timezone.now().strftime("%Y-%m-%d %H:%M")}')
        date_cell.alignment = Alignment(horizontal='center')

        # Headers
        header_fill = PatternFill(start_color='1a73e8', end_color='1a73e8', fill_type='solid')
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin'),
        )

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border

        ws.row_dimensions[4].height = 24

        # Data
        for row_idx, row_data in enumerate(data, 5):
            fill_color = 'F8F9FA' if row_idx % 2 == 1 else 'FFFFFF'
            row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')
                cell.fill = row_fill
                if isinstance(value, (int, float, Decimal)):
                    cell.number_format = '#,##0.00'

        # تعديل عرض الأعمدة
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18

        # Response
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        if not filename.endswith('.xlsx'):
            filename += '.xlsx'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        wb.save(response)
        return response


class PDFExporter:
    """تصدير PDF"""

    @classmethod
    def export(cls, template_name, context, filename='report.pdf'):
        """تصدير HTML إلى PDF عبر xhtml2pdf"""
        from xhtml2pdf import pisa

        html = render_to_string(template_name, context)
        response = HttpResponse(content_type='application/pdf')
        if not filename.endswith('.pdf'):
            filename += '.pdf'
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        pisa_status = pisa.CreatePDF(html, dest=response, encoding='utf-8')

        if pisa_status.err:
            return HttpResponse('خطأ في إنشاء PDF', status=500)

        return response
