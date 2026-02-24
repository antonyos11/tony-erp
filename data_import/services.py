# -*- coding: utf-8 -*-
"""
خدمات استيراد البيانات من Excel
"""
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from django.db import transaction
from django.utils import timezone
from django.contrib.auth.models import User
from decimal import Decimal, InvalidOperation
from datetime import datetime
from typing import Dict, List, Any, Optional, Tuple, Sequence
import logging

logger = logging.getLogger(__name__)


class ExcelImportService:
    """خدمة استيراد البيانات من Excel"""
    
    # تعريف أعمدة كل موديول
    COLUMN_MAPPINGS = {
        'categories': {
            'name': ['الاسم', 'اسم الفئة', 'Category Name', 'Name'],
            'description': ['الوصف', 'Description'],
            'parent': ['الفئة الأم', 'Parent Category', 'Parent'],
            'is_active': ['نشط', 'Active', 'Is Active'],
        },
        'products': {
            'sku': ['الكود', 'كود المنتج', 'SKU', 'Code', 'رقم الصنف'],
            'name': ['الاسم', 'اسم المنتج', 'Product Name', 'Name', 'اسم الصنف'],
            'description': ['الوصف', 'Description'],
            'category': ['الفئة', 'Category', 'التصنيف'],
            'price': ['السعر', 'سعر البيع', 'Price', 'Selling Price'],
            'cost': ['التكلفة', 'سعر التكلفة', 'Cost', 'Cost Price'],
            'min_stock': ['الحد الأدنى', 'Min Stock', 'Minimum Stock'],
            'unit': ['الوحدة', 'Unit', 'UOM'],
            'barcode': ['الباركود', 'Barcode'],
            'is_active': ['نشط', 'Active'],
        },
        'customers': {
            'name': ['الاسم', 'اسم العميل', 'Customer Name', 'Name'],
            'email': ['البريد الإلكتروني', 'Email', 'البريد'],
            'phone': ['الهاتف', 'رقم الهاتف', 'Phone', 'Mobile'],
            'address': ['العنوان', 'Address'],
            'is_key_account': ['عميل رئيسي', 'Key Account', 'VIP'],
        },
        'suppliers': {
            'name': ['الاسم', 'اسم المورد', 'Supplier Name', 'Name'],
            'email': ['البريد الإلكتروني', 'Email'],
            'phone': ['الهاتف', 'Phone'],
            'address': ['العنوان', 'Address'],
            'supply_type': ['نوع التوريد', 'Supply Type'],
            'raw_material': ['المادة الخام', 'Raw Material'],
        },
        'employees': {
            'employee_id': ['رقم الموظف', 'Employee ID', 'Code'],
            'first_name': ['الاسم الأول', 'First Name'],
            'last_name': ['اسم العائلة', 'Last Name', 'Family Name'],
            'arabic_name': ['الاسم بالعربية', 'Arabic Name', 'الاسم الكامل'],
            'national_id': ['رقم الهوية', 'National ID', 'الرقم القومي'],
            'gender': ['الجنس', 'Gender'],
            'birth_date': ['تاريخ الميلاد', 'Birth Date', 'DOB'],
            'phone': ['الهاتف', 'Phone'],
            'email': ['البريد', 'Email'],
            'address': ['العنوان', 'Address'],
            'department': ['القسم', 'Department'],
            'position': ['المنصب', 'Position', 'الوظيفة'],
            'hire_date': ['تاريخ التعيين', 'Hire Date', 'Join Date'],
            'basic_salary': ['الراتب الأساسي', 'Basic Salary', 'Salary'],
        },
        'departments': {
            'code': ['الكود', 'Code', 'Department Code'],
            'name': ['الاسم', 'اسم القسم', 'Department Name', 'Name'],
            'description': ['الوصف', 'Description'],
            'budget': ['الميزانية', 'Budget'],
        },
        'accounts': {
            'code': ['الكود', 'رقم الحساب', 'Account Code', 'Code'],
            'name': ['الاسم', 'اسم الحساب', 'Account Name', 'Name'],
            'account_type': ['النوع', 'نوع الحساب', 'Account Type', 'Type'],
            'parent': ['الحساب الأب', 'Parent Account', 'Parent'],
            'description': ['الوصف', 'Description'],
        },
        'locations': {
            'code': ['الكود', 'Code', 'Location Code'],
            'name': ['الاسم', 'اسم المخزن', 'Location Name', 'Warehouse Name'],
            'address': ['العنوان', 'Address'],
            'location_type': ['النوع', 'Type', 'Location Type'],
        },
        'opening_balances': {
            'sku': ['كود المنتج', 'SKU', 'Product Code'],
            'location': ['المخزن', 'Location', 'Warehouse'],
            'quantity': ['الكمية', 'Quantity', 'Qty'],
            'unit_cost': ['تكلفة الوحدة', 'Unit Cost', 'Cost'],
        },
    }
    
    def __init__(self, session=None):
        self.session = session
        self.errors = []
    
    def find_column(self, headers: Sequence[Any], field: str, module: str) -> Optional[int]:
        """البحث عن عمود في الهيدرز"""
        possible_names = self.COLUMN_MAPPINGS.get(module, {}).get(field, [])
        for idx, header in enumerate(headers):
            if header and str(header).strip() in possible_names:
                return idx
        return None
    
    def parse_value(self, value: Any, field_type: str = 'string') -> Any:
        """تحويل القيمة حسب النوع"""
        if value is None or (isinstance(value, str) and not value.strip()):
            return None
        
        if field_type == 'string':
            return str(value).strip()
        elif field_type == 'decimal':
            try:
                return Decimal(str(value).replace(',', ''))
            except (InvalidOperation, ValueError):
                return Decimal('0')
        elif field_type == 'integer':
            try:
                return int(float(str(value)))
            except (ValueError, TypeError):
                return 0
        elif field_type == 'boolean':
            if isinstance(value, bool):
                return value
            str_val = str(value).lower().strip()
            return str_val in ('yes', 'نعم', 'true', '1', 'active', 'نشط', 'صحيح')
        elif field_type == 'date':
            if isinstance(value, datetime):
                return value.date()
            try:
                # محاولة تحويل التاريخ من عدة صيغ
                for fmt in ['%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d']:
                    try:
                        return datetime.strptime(str(value), fmt).date()
                    except ValueError:
                        continue
            except Exception:
                pass
            return None
        
        return value
    
    def log_error(self, row_number: int, column: str, error_type: str, message: str, row_data: Optional[Dict[str, Any]] = None):
        """تسجيل خطأ"""
        from .models import ImportError
        
        error: Dict[str, Any] = {
            'row_number': row_number,
            'column_name': column,
            'error_type': error_type,
            'error_message': message,
            'row_data': row_data if row_data is not None else {},
        }
        self.errors.append(error)
        
        if self.session:
            ImportError.objects.create(
                session=self.session,
                **error
            )
    
    def import_categories(self, file_path: str) -> Tuple[int, int]:
        """استيراد الفئات"""
        from inventory.models import Category
        
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        if ws is None:
            return 0, 0
        
        headers = [cell.value for cell in ws[1]]
        
        # البحث عن الأعمدة
        name_col = self.find_column(headers, 'name', 'categories')
        desc_col = self.find_column(headers, 'description', 'categories')
        parent_col = self.find_column(headers, 'parent', 'categories')
        active_col = self.find_column(headers, 'is_active', 'categories')
        
        if name_col is None:
            self.log_error(1, 'name', 'missing_column', 'عمود اسم الفئة مطلوب')
            return 0, 1
        
        success_count = 0
        error_count = 0
        
        # المرور الأول: إنشاء جميع الفئات بدون الأب
        categories_map = {}
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            try:
                name = self.parse_value(row[name_col].value)
                if not name:
                    continue
                
                description = self.parse_value(row[desc_col].value) if desc_col else ''
                is_active = self.parse_value(row[active_col].value, 'boolean') if active_col else True
                
                category, created = Category.objects.update_or_create(
                    name=name,
                    defaults={
                        'description': description or '',
                        'is_active': is_active if is_active is not None else True,
                    }
                )
                categories_map[name] = category
                success_count += 1
                
            except Exception as e:
                self.log_error(row_idx, '', 'import_error', str(e))
                error_count += 1
        
        # المرور الثاني: تحديث الفئات الأم
        if parent_col is not None:
            for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                try:
                    name = self.parse_value(row[name_col].value)
                    parent_name = self.parse_value(row[parent_col].value)
                    
                    if name and parent_name and name in categories_map:
                        parent_cat = categories_map.get(parent_name) or Category.objects.filter(name=parent_name).first()
                        if parent_cat:
                            cat = categories_map[name]
                            cat.parent = parent_cat
                            cat.save()
                except Exception:
                    pass
        
        wb.close()
        return success_count, error_count
    
    def import_products(self, file_path: str) -> Tuple[int, int]:
        """استيراد المنتجات"""
        from inventory.models import Product, Category, generate_unique_barcode
        
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        if ws is None:
            return 0, 0
        
        headers = [cell.value for cell in ws[1]]
        
        # البحث عن الأعمدة
        col_indices = {}
        for field in ['sku', 'name', 'description', 'category', 'price', 'cost', 'min_stock', 'unit', 'barcode', 'is_active']:
            col_indices[field] = self.find_column(headers, field, 'products')
        
        if col_indices['name'] is None:
            self.log_error(1, 'name', 'missing_column', 'عمود اسم المنتج مطلوب')
            return 0, 1
        
        success_count = 0
        error_count = 0
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            try:
                name = self.parse_value(row[col_indices['name']].value)
                if not name:
                    continue
                
                # توليد SKU إذا لم يوجد
                sku = None
                if col_indices['sku'] is not None:
                    sku = self.parse_value(row[col_indices['sku']].value)
                if not sku:
                    from core.sequence_utils import next_sequence, format_code
                    seq = next_sequence('product')
                    sku = format_code('PRD-', seq)
                
                # البحث عن الفئة
                category = None
                if col_indices['category'] is not None:
                    cat_name = self.parse_value(row[col_indices['category']].value)
                    if cat_name:
                        category = Category.objects.filter(name=cat_name).first()
                
                # الباركود
                barcode = None
                if col_indices['barcode'] is not None:
                    barcode = self.parse_value(row[col_indices['barcode']].value)
                if not barcode:
                    barcode = generate_unique_barcode()
                
                product_data = {
                    'name': name,
                    'description': self.parse_value(row[col_indices['description']].value) if col_indices['description'] else '',
                    'category': category,
                    'price': self.parse_value(row[col_indices['price']].value, 'decimal') if col_indices['price'] else Decimal('0'),
                    'cost': self.parse_value(row[col_indices['cost']].value, 'decimal') if col_indices['cost'] else Decimal('0'),
                    'min_stock': self.parse_value(row[col_indices['min_stock']].value, 'integer') if col_indices['min_stock'] else 0,
                    'barcode': barcode,
                }
                
                if col_indices['unit'] is not None:
                    unit = self.parse_value(row[col_indices['unit']].value)
                    if unit:
                        # تحويل اسم الوحدة إلى كود
                        unit_map = {'وحدة': 'unit', 'كيلو': 'kg', 'كيلوجرام': 'kg', 'جرام': 'g', 'متر': 'm', 'لتر': 'l'}
                        product_data['unit'] = unit_map.get(unit, 'unit')
                
                product, created = Product.objects.update_or_create(
                    sku=sku,
                    defaults=product_data
                )
                success_count += 1
                
            except Exception as e:
                self.log_error(row_idx, '', 'import_error', str(e))
                error_count += 1
        
        wb.close()
        return success_count, error_count
    
    def import_customers(self, file_path: str) -> Tuple[int, int]:
        """استيراد العملاء"""
        from partners.models import Customer, Partner
        
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        if ws is None:
            return 0, 0
        
        headers = [cell.value for cell in ws[1]]
        
        col_indices = {}
        for field in ['name', 'email', 'phone', 'address', 'is_key_account']:
            col_indices[field] = self.find_column(headers, field, 'customers')
        
        if col_indices['name'] is None:
            self.log_error(1, 'name', 'missing_column', 'عمود اسم العميل مطلوب')
            return 0, 1
        
        success_count = 0
        error_count = 0
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            try:
                name = self.parse_value(row[col_indices['name']].value)
                if not name:
                    continue
                
                customer_data = {
                    'email': self.parse_value(row[col_indices['email']].value) if col_indices['email'] else '',
                    'phone': self.parse_value(row[col_indices['phone']].value) if col_indices['phone'] else '',
                    'address': self.parse_value(row[col_indices['address']].value) if col_indices['address'] else '',
                    'is_key_account': self.parse_value(row[col_indices['is_key_account']].value, 'boolean') if col_indices['is_key_account'] else False,
                }
                
                customer, created = Customer.objects.update_or_create(
                    name=name,
                    defaults=customer_data
                )
                success_count += 1
                
            except Exception as e:
                self.log_error(row_idx, '', 'import_error', str(e))
                error_count += 1
        
        wb.close()
        return success_count, error_count
    
    def import_suppliers(self, file_path: str) -> Tuple[int, int]:
        """استيراد الموردين"""
        from partners.models import Supplier, Partner
        
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        if ws is None:
            return 0, 0
        
        headers = [cell.value for cell in ws[1]]
        
        col_indices = {}
        for field in ['name', 'email', 'phone', 'address', 'supply_type', 'raw_material']:
            col_indices[field] = self.find_column(headers, field, 'suppliers')
        
        if col_indices['name'] is None:
            self.log_error(1, 'name', 'missing_column', 'عمود اسم المورد مطلوب')
            return 0, 1
        
        success_count = 0
        error_count = 0
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            try:
                name = self.parse_value(row[col_indices['name']].value)
                if not name:
                    continue
                
                # تحويل نوع التوريد
                supply_type = ''
                if col_indices['supply_type'] is not None:
                    st = self.parse_value(row[col_indices['supply_type']].value)
                    if st:
                        supply_map = {'خامات': 'raw_materials', 'مواد تعبئة': 'packaging', 'خدمات': 'services', 'قطع غيار': 'spare_parts'}
                        supply_type = supply_map.get(st, 'other')
                
                supplier_data = {
                    'email': self.parse_value(row[col_indices['email']].value) if col_indices['email'] else '',
                    'phone': self.parse_value(row[col_indices['phone']].value) if col_indices['phone'] else '',
                    'address': self.parse_value(row[col_indices['address']].value) if col_indices['address'] else '',
                    'supply_type': supply_type,
                    'raw_material': self.parse_value(row[col_indices['raw_material']].value) if col_indices['raw_material'] else '',
                }
                
                supplier, created = Supplier.objects.update_or_create(
                    name=name,
                    defaults=supplier_data
                )
                success_count += 1
                
            except Exception as e:
                self.log_error(row_idx, '', 'import_error', str(e))
                error_count += 1
        
        wb.close()
        return success_count, error_count
    
    def import_departments(self, file_path: str) -> Tuple[int, int]:
        """استيراد الأقسام"""
        from hr.models import Department
        
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        if ws is None:
            return 0, 0
        
        headers = [cell.value for cell in ws[1]]
        
        col_indices = {}
        for field in ['code', 'name', 'description', 'budget']:
            col_indices[field] = self.find_column(headers, field, 'departments')
        
        if col_indices['name'] is None:
            self.log_error(1, 'name', 'missing_column', 'عمود اسم القسم مطلوب')
            return 0, 1
        
        success_count = 0
        error_count = 0
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            try:
                name = self.parse_value(row[col_indices['name']].value)
                if not name:
                    continue
                
                code = None
                if col_indices['code'] is not None:
                    code = self.parse_value(row[col_indices['code']].value)
                if not code:
                    code = f"DEP-{row_idx:04d}"
                
                dept_data = {
                    'name': name,
                    'description': self.parse_value(row[col_indices['description']].value) if col_indices['description'] else '',
                    'budget': self.parse_value(row[col_indices['budget']].value, 'decimal') if col_indices['budget'] else Decimal('0'),
                }
                
                dept, created = Department.objects.update_or_create(
                    code=code,
                    defaults=dept_data
                )
                success_count += 1
                
            except Exception as e:
                self.log_error(row_idx, '', 'import_error', str(e))
                error_count += 1
        
        wb.close()
        return success_count, error_count
    
    def import_employees(self, file_path: str) -> Tuple[int, int]:
        """استيراد الموظفين"""
        from hr.models import Employee, Department, JobPosition
        from django.contrib.auth.models import User
        
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        if ws is None:
            return 0, 0
        
        headers = [cell.value for cell in ws[1]]
        
        col_indices = {}
        for field in ['employee_id', 'first_name', 'last_name', 'arabic_name', 'national_id', 
                      'gender', 'birth_date', 'phone', 'email', 'address', 'department', 
                      'position', 'hire_date', 'basic_salary']:
            col_indices[field] = self.find_column(headers, field, 'employees')
        
        if col_indices['arabic_name'] is None and col_indices['first_name'] is None:
            self.log_error(1, 'name', 'missing_column', 'عمود اسم الموظف مطلوب')
            return 0, 1
        
        success_count = 0
        error_count = 0
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            try:
                # الحصول على الاسم
                arabic_name = None
                if col_indices['arabic_name'] is not None:
                    arabic_name = self.parse_value(row[col_indices['arabic_name']].value)
                
                first_name = self.parse_value(row[col_indices['first_name']].value) if col_indices['first_name'] else ''
                last_name = self.parse_value(row[col_indices['last_name']].value) if col_indices['last_name'] else ''
                
                if not arabic_name and not first_name:
                    continue
                
                if not arabic_name:
                    arabic_name = f"{first_name} {last_name}".strip()
                if not first_name:
                    parts = arabic_name.split()
                    first_name = parts[0] if parts else 'Employee'
                    last_name = parts[-1] if len(parts) > 1 else ''
                
                # رقم الموظف
                employee_id = None
                if col_indices['employee_id'] is not None:
                    employee_id = self.parse_value(row[col_indices['employee_id']].value)
                if not employee_id:
                    employee_id = f"EMP-{row_idx:05d}"
                
                # الرقم القومي
                national_id = None
                if col_indices['national_id'] is not None:
                    national_id = self.parse_value(row[col_indices['national_id']].value)
                if not national_id:
                    national_id = f"ID-{row_idx:010d}"
                
                # القسم
                department = None
                if col_indices['department'] is not None:
                    dept_name = self.parse_value(row[col_indices['department']].value)
                    if dept_name:
                        department = Department.objects.filter(name=dept_name).first()
                if not department:
                    department, _ = Department.objects.get_or_create(
                        code='DEFAULT',
                        defaults={'name': 'القسم العام'}
                    )
                
                # المنصب
                position = None
                if col_indices['position'] is not None:
                    pos_name = self.parse_value(row[col_indices['position']].value)
                    if pos_name:
                        position = JobPosition.objects.filter(title=pos_name).first()
                if not position:
                    position, _ = JobPosition.objects.get_or_create(
                        code='DEFAULT',
                        defaults={
                            'title': 'موظف',
                            'department': department,
                            'description': 'منصب افتراضي',
                            'requirements': '-'
                        }
                    )
                
                # إنشاء مستخدم
                username = f"emp_{employee_id.lower().replace('-', '_')}"
                user, user_created = User.objects.get_or_create(
                    username=username,
                    defaults={
                        'first_name': first_name,
                        'last_name': last_name,
                        'email': self.parse_value(row[col_indices['email']].value) if col_indices['email'] else '',
                    }
                )
                if user_created:
                    user.set_password('changeme123')
                    user.save()
                
                # تاريخ الميلاد
                birth_date = None
                if col_indices['birth_date'] is not None:
                    birth_date = self.parse_value(row[col_indices['birth_date']].value, 'date')
                if not birth_date:
                    birth_date = timezone.now().date().replace(year=timezone.now().year - 30)
                
                # تاريخ التعيين
                hire_date = None
                if col_indices['hire_date'] is not None:
                    hire_date = self.parse_value(row[col_indices['hire_date']].value, 'date')
                if not hire_date:
                    hire_date = timezone.now().date()
                
                # الجنس
                gender = 'M'
                if col_indices['gender'] is not None:
                    g = self.parse_value(row[col_indices['gender']].value)
                    if g:
                        gender = 'F' if g.lower() in ('f', 'female', 'أنثى', 'انثى') else 'M'
                
                emp_data = {
                    'user': user,
                    'first_name': first_name,
                    'last_name': last_name,
                    'arabic_name': arabic_name,
                    'national_id': national_id,
                    'gender': gender,
                    'birth_date': birth_date,
                    'marital_status': 'single',
                    'phone': self.parse_value(row[col_indices['phone']].value) if col_indices['phone'] else '',
                    'email': self.parse_value(row[col_indices['email']].value) if col_indices['email'] else '',
                    'address': self.parse_value(row[col_indices['address']].value) if col_indices['address'] else '',
                    'emergency_contact_name': '',
                    'emergency_contact_phone': '',
                    'department': department,
                    'position': position,
                    'hire_date': hire_date,
                    'basic_salary': self.parse_value(row[col_indices['basic_salary']].value, 'decimal') if col_indices['basic_salary'] else Decimal('0'),
                }
                
                employee, created = Employee.objects.update_or_create(
                    employee_id=employee_id,
                    defaults=emp_data
                )
                success_count += 1
                
            except Exception as e:
                self.log_error(row_idx, '', 'import_error', str(e))
                error_count += 1
        
        wb.close()
        return success_count, error_count
    
    def import_locations(self, file_path: str) -> Tuple[int, int]:
        """استيراد المخازن/المواقع"""
        from inventory.models import Location
        
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        if ws is None:
            return 0, 0
        
        headers = [cell.value for cell in ws[1]]
        
        col_indices = {}
        for field in ['code', 'name', 'address', 'location_type']:
            col_indices[field] = self.find_column(headers, field, 'locations')
        
        if col_indices['name'] is None:
            self.log_error(1, 'name', 'missing_column', 'عمود اسم المخزن مطلوب')
            return 0, 1
        
        success_count = 0
        error_count = 0
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            try:
                name = self.parse_value(row[col_indices['name']].value)
                if not name:
                    continue
                
                code = None
                if col_indices['code'] is not None:
                    code = self.parse_value(row[col_indices['code']].value)
                if not code:
                    code = f"LOC-{row_idx:04d}"
                
                loc_data = {
                    'name': name,
                    'address': self.parse_value(row[col_indices['address']].value) if col_indices['address'] else '',
                }
                
                location, created = Location.objects.update_or_create(
                    code=code,
                    defaults=loc_data
                )
                success_count += 1
                
            except Exception as e:
                self.log_error(row_idx, '', 'import_error', str(e))
                error_count += 1
        
        wb.close()
        return success_count, error_count
    
    def import_accounts(self, file_path: str) -> Tuple[int, int]:
        """استيراد الحسابات المحاسبية"""
        from accounting.models import Account
        
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        if ws is None:
            return 0, 0
        
        headers = [cell.value for cell in ws[1]]
        
        col_indices = {}
        for field in ['code', 'name', 'account_type', 'parent', 'description']:
            col_indices[field] = self.find_column(headers, field, 'accounts')
        
        if col_indices['name'] is None or col_indices['code'] is None:
            self.log_error(1, 'name', 'missing_column', 'عمود اسم الحساب وكوده مطلوبان')
            return 0, 1
        
        success_count = 0
        error_count = 0
        accounts_map = {}
        
        # المرور الأول: إنشاء جميع الحسابات
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            try:
                code = self.parse_value(row[col_indices['code']].value)
                name = self.parse_value(row[col_indices['name']].value)
                if not name or not code:
                    continue
                
                # نوع الحساب
                account_type = 'asset'
                if col_indices['account_type'] is not None:
                    at = self.parse_value(row[col_indices['account_type']].value)
                    if at:
                        type_map = {
                            'أصول': 'asset', 'أصل': 'asset', 'asset': 'asset', 'assets': 'asset',
                            'التزامات': 'liability', 'خصوم': 'liability', 'liability': 'liability',
                            'حقوق ملكية': 'equity', 'equity': 'equity',
                            'إيرادات': 'revenue', 'إيراد': 'revenue', 'revenue': 'revenue', 'income': 'revenue',
                            'مصروفات': 'expense', 'مصروف': 'expense', 'expense': 'expense', 'expenses': 'expense',
                        }
                        account_type = type_map.get(at.lower(), 'asset')
                
                acc_data = {
                    'name': name,
                    'account_type': account_type,
                    'description': self.parse_value(row[col_indices['description']].value) if col_indices['description'] else '',
                }
                
                account, created = Account.objects.update_or_create(
                    code=code,
                    defaults=acc_data
                )
                accounts_map[code] = account
                success_count += 1
                
            except Exception as e:
                self.log_error(row_idx, '', 'import_error', str(e))
                error_count += 1
        
        # المرور الثاني: تحديث الحسابات الأم
        if col_indices['parent'] is not None:
            for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                try:
                    code = self.parse_value(row[col_indices['code']].value)
                    parent_code = self.parse_value(row[col_indices['parent']].value)
                    
                    if code and parent_code and code in accounts_map:
                        parent_acc = accounts_map.get(parent_code) or Account.objects.filter(code=parent_code).first()
                        if parent_acc:
                            acc = accounts_map[code]
                            acc.parent = parent_acc
                            acc.save()
                except Exception:
                    pass
        
        wb.close()
        return success_count, error_count
    
    def import_opening_balances(self, file_path: str) -> Tuple[int, int]:
        """استيراد الأرصدة الافتتاحية للمخزون"""
        from inventory.models import Product, Location, Stock
        
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        if ws is None:
            return 0, 0
        
        headers = [cell.value for cell in ws[1]]
        
        col_indices = {}
        for field in ['sku', 'location', 'quantity', 'unit_cost']:
            col_indices[field] = self.find_column(headers, field, 'opening_balances')
        
        if col_indices['sku'] is None or col_indices['quantity'] is None:
            self.log_error(1, 'sku', 'missing_column', 'عمود كود المنتج والكمية مطلوبان')
            return 0, 1
        
        success_count = 0
        error_count = 0
        
        # الحصول على المخزن الافتراضي
        default_location = Location.objects.first()
        if not default_location:
            default_location = Location.objects.create(code='MAIN', name='المخزن الرئيسي')
        
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            try:
                sku = self.parse_value(row[col_indices['sku']].value)
                quantity = self.parse_value(row[col_indices['quantity']].value, 'decimal')
                
                if not sku or quantity is None:
                    continue
                
                product = Product.objects.filter(sku=sku).first()
                if not product:
                    self.log_error(row_idx, 'sku', 'not_found', f'المنتج غير موجود: {sku}')
                    error_count += 1
                    continue
                
                # المخزن
                location = default_location
                if col_indices['location'] is not None:
                    loc_name = self.parse_value(row[col_indices['location']].value)
                    if loc_name:
                        loc = Location.objects.filter(name=loc_name).first() or Location.objects.filter(code=loc_name).first()
                        if loc:
                            location = loc
                
                # تكلفة الوحدة
                unit_cost = Decimal('0')
                if col_indices['unit_cost'] is not None:
                    unit_cost = self.parse_value(row[col_indices['unit_cost']].value, 'decimal') or Decimal('0')
                
                stock, created = Stock.objects.update_or_create(
                    product=product,
                    location=location,
                    defaults={
                        'quantity': quantity,
                        'unit_cost': unit_cost,
                    }
                )
                success_count += 1
                
            except Exception as e:
                self.log_error(row_idx, '', 'import_error', str(e))
                error_count += 1
        
        wb.close()
        return success_count, error_count


class ExcelTemplateGenerator:
    """مولد قوالب Excel"""
    
    TEMPLATES = {
        'categories': {
            'name': 'قالب_الفئات.xlsx',
            'headers': ['اسم الفئة', 'الوصف', 'الفئة الأم', 'نشط'],
            'sample_data': [
                ['إلكترونيات', 'أجهزة إلكترونية ومستلزماتها', '', 'نعم'],
                ['هواتف', 'هواتف ذكية وملحقاتها', 'إلكترونيات', 'نعم'],
                ['لابتوب', 'أجهزة كمبيوتر محمولة', 'إلكترونيات', 'نعم'],
            ]
        },
        'products': {
            'name': 'قالب_المنتجات.xlsx',
            'headers': ['كود المنتج', 'اسم المنتج', 'الوصف', 'الفئة', 'سعر البيع', 'سعر التكلفة', 'الحد الأدنى', 'الوحدة', 'الباركود'],
            'sample_data': [
                ['PRD-001', 'آيفون 15', 'هاتف ذكي من آبل', 'هواتف', 45000, 40000, 5, 'وحدة', '123456789012'],
                ['PRD-002', 'سامسونج S24', 'هاتف ذكي من سامسونج', 'هواتف', 35000, 30000, 5, 'وحدة', '123456789013'],
            ]
        },
        'customers': {
            'name': 'قالب_العملاء.xlsx',
            'headers': ['اسم العميل', 'البريد الإلكتروني', 'رقم الهاتف', 'العنوان', 'عميل رئيسي'],
            'sample_data': [
                ['شركة الأمل', 'info@alamal.com', '01234567890', 'القاهرة - مصر', 'نعم'],
                ['محمد أحمد', 'mohamed@email.com', '01098765432', 'الإسكندرية - مصر', 'لا'],
            ]
        },
        'suppliers': {
            'name': 'قالب_الموردين.xlsx',
            'headers': ['اسم المورد', 'البريد الإلكتروني', 'رقم الهاتف', 'العنوان', 'نوع التوريد', 'المادة الخام'],
            'sample_data': [
                ['شركة التوريدات', 'supply@co.com', '01234567890', 'القاهرة', 'خامات', 'مواد خام'],
                ['مصنع التعبئة', 'pack@factory.com', '01098765432', 'الجيزة', 'مواد تعبئة', 'عبوات بلاستيكية'],
            ]
        },
        'employees': {
            'name': 'قالب_الموظفين.xlsx',
            'headers': ['رقم الموظف', 'الاسم بالعربية', 'الاسم الأول', 'اسم العائلة', 'رقم الهوية', 
                       'الجنس', 'تاريخ الميلاد', 'الهاتف', 'البريد', 'العنوان', 
                       'القسم', 'الوظيفة', 'تاريخ التعيين', 'الراتب الأساسي'],
            'sample_data': [
                ['EMP-001', 'أحمد محمد علي', 'أحمد', 'علي', '29901011234567', 
                 'ذكر', '1990-01-01', '01234567890', 'ahmed@company.com', 'القاهرة',
                 'المبيعات', 'مندوب مبيعات', '2023-01-15', 8000],
            ]
        },
        'departments': {
            'name': 'قالب_الأقسام.xlsx',
            'headers': ['كود القسم', 'اسم القسم', 'الوصف', 'الميزانية'],
            'sample_data': [
                ['SALES', 'المبيعات', 'قسم المبيعات والتسويق', 500000],
                ['HR', 'الموارد البشرية', 'قسم شؤون الموظفين', 200000],
                ['ACCT', 'المحاسبة', 'قسم الحسابات والمالية', 300000],
            ]
        },
        'accounts': {
            'name': 'قالب_الحسابات.xlsx',
            'headers': ['رقم الحساب', 'اسم الحساب', 'نوع الحساب', 'الحساب الأب', 'الوصف'],
            'sample_data': [
                ['1', 'الأصول', 'أصول', '', 'إجمالي الأصول'],
                ['11', 'الأصول المتداولة', 'أصول', '1', 'الأصول قصيرة الأجل'],
                ['111', 'النقدية والبنوك', 'أصول', '11', 'الأرصدة النقدية'],
                ['2', 'الخصوم', 'التزامات', '', 'إجمالي الالتزامات'],
                ['4', 'الإيرادات', 'إيرادات', '', 'إجمالي الإيرادات'],
                ['5', 'المصروفات', 'مصروفات', '', 'إجمالي المصروفات'],
            ]
        },
        'locations': {
            'name': 'قالب_المخازن.xlsx',
            'headers': ['كود المخزن', 'اسم المخزن', 'العنوان'],
            'sample_data': [
                ['MAIN', 'المخزن الرئيسي', 'القاهرة - المنطقة الصناعية'],
                ['BRANCH1', 'فرع الإسكندرية', 'الإسكندرية - شارع الحرية'],
            ]
        },
        'opening_balances': {
            'name': 'قالب_الأرصدة_الافتتاحية.xlsx',
            'headers': ['كود المنتج', 'المخزن', 'الكمية', 'تكلفة الوحدة'],
            'sample_data': [
                ['PRD-001', 'المخزن الرئيسي', 100, 40000],
                ['PRD-002', 'المخزن الرئيسي', 50, 30000],
            ]
        },
    }
    
    @classmethod
    def generate_template(cls, module: str) -> Workbook:
        """إنشاء قالب Excel لموديول معين"""
        template = cls.TEMPLATES.get(module)
        if not template:
            raise ValueError(f"قالب غير موجود: {module}")
        
        wb = Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet()
        
        ws.title = 'البيانات'
        
        # تنسيق الهيدر
        header_font = Font(bold=True, color='FFFFFF', size=12)
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # كتابة الهيدرز
        for col_idx, header in enumerate(template['headers'], start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
            
            # تعيين عرض العمود
            ws.column_dimensions[get_column_letter(col_idx)].width = max(15, len(header) + 5)
        
        # كتابة البيانات النموذجية
        data_alignment = Alignment(horizontal='center', vertical='center')
        
        for row_idx, row_data in enumerate(template['sample_data'], start=2):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = data_alignment
                cell.border = thin_border
        
        # تجميد الصف الأول
        ws.freeze_panes = 'A2'
        
        # إضافة ورقة التعليمات
        instructions_ws = wb.create_sheet('تعليمات')
        instructions = [
            ('تعليمات الاستخدام', ''),
            ('', ''),
            ('1. احذف البيانات النموذجية واستبدلها ببياناتك', ''),
            ('2. لا تغير أسماء الأعمدة في الصف الأول', ''),
            ('3. تأكد من عدم وجود صفوف فارغة بين البيانات', ''),
            ('4. احفظ الملف بصيغة xlsx', ''),
            ('', ''),
            ('ملاحظات:', ''),
            ('- الأعمدة المطلوبة موضحة في الصف الأول', ''),
            ('- يمكنك ترك الأعمدة الاختيارية فارغة', ''),
        ]
        
        for row_idx, (text, _) in enumerate(instructions, start=1):
            cell = instructions_ws.cell(row=row_idx, column=1, value=text)
            if row_idx == 1:
                cell.font = Font(bold=True, size=14)
            elif text.startswith('ملاحظات'):
                cell.font = Font(bold=True)
        
        instructions_ws.column_dimensions['A'].width = 50
        
        return wb
    
    @classmethod
    def get_all_templates_zip(cls):
        """إنشاء ملف zip يحتوي على جميع القوالب"""
        import io
        import zipfile
        
        zip_buffer = io.BytesIO()
        
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            for module, template in cls.TEMPLATES.items():
                wb = cls.generate_template(module)
                excel_buffer = io.BytesIO()
                wb.save(excel_buffer)
                excel_buffer.seek(0)
                zip_file.writestr(template['name'], excel_buffer.read())
                wb.close()
        
        zip_buffer.seek(0)
        return zip_buffer
