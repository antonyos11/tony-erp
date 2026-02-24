"""
Views لمنشئ التقارير المرئي
"""

import json
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_http_methods
from django.contrib import messages
from django.utils import timezone
from django.core.paginator import Paginator

from .models import (
    DataSource, Report, ReportWidget, ReportTemplate,
    ScheduledReport, ReportExecution
)
from .services import DataSourceService, ReportExecutionService, ChartDataService


@login_required
def dashboard(request):
    """لوحة التقارير"""
    my_reports = Report.objects.filter(created_by=request.user).order_by('-updated_at')[:10]
    shared_reports = Report.objects.filter(shared_with=request.user).order_by('-updated_at')[:10]
    recent_executions = ReportExecution.objects.filter(
        executed_by=request.user
    ).order_by('-started_at')[:10]
    
    templates = ReportTemplate.objects.filter(is_active=True)[:6]
    
    stats = {
        'total_reports': Report.objects.filter(created_by=request.user).count(),
        'scheduled': ScheduledReport.objects.filter(report__created_by=request.user, is_active=True).count(),
        'executions_today': ReportExecution.objects.filter(
            executed_by=request.user,
            started_at__date=timezone.now().date()
        ).count(),
    }
    
    return render(request, 'report_builder/dashboard.html', {
        'my_reports': my_reports,
        'shared_reports': shared_reports,
        'recent_executions': recent_executions,
        'templates': templates,
        'stats': stats,
    })


@login_required
def report_list(request):
    """قائمة التقارير"""
    reports = Report.objects.filter(
        created_by=request.user
    ).order_by('-updated_at')
    
    # البحث
    search = request.GET.get('q', '')
    if search:
        reports = reports.filter(name__icontains=search)
    
    # التصفية
    report_type = request.GET.get('type')
    if report_type:
        reports = reports.filter(report_type=report_type)
    
    paginator = Paginator(reports, 20)
    page = request.GET.get('page', 1)
    reports = paginator.get_page(page)
    
    return render(request, 'report_builder/list.html', {
        'reports': reports,
    })


@login_required
def create_report(request):
    """إنشاء تقرير جديد"""
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            
            # إنشاء مصدر بيانات إذا لزم الأمر
            data_source_id = data.get('data_source_id')
            if not data_source_id and data.get('content_type_id'):
                data_source = DataSource.objects.create(
                    name=f"مصدر {data.get('name')}",
                    source_type='model',
                    content_type_id=data.get('content_type_id'),
                    created_by=request.user
                )
                data_source_id = data_source.id
            
            report = Report.objects.create(
                name=data.get('name'),
                description=data.get('description', ''),
                report_type=data.get('report_type', 'table'),
                data_source_id=data_source_id,
                selected_columns=data.get('columns', []),
                filters=data.get('filters', []),
                sorting=data.get('sorting', []),
                grouping=data.get('grouping', []),
                calculations=data.get('calculations', []),
                created_by=request.user
            )
            
            return JsonResponse({
                'success': True,
                'report_id': str(report.uuid)
            })
            
        except Exception as e:
            return JsonResponse({'success': False, 'error': str(e)})
    
    # الحصول على النماذج المتاحة
    available_models = DataSourceService.get_available_models()
    data_sources = DataSource.objects.filter(created_by=request.user)
    templates = ReportTemplate.objects.filter(is_active=True)
    
    return render(request, 'report_builder/create.html', {
        'available_models': available_models,
        'data_sources': data_sources,
        'templates': templates,
    })


@login_required
def report_builder(request, uuid):
    """منشئ التقرير المرئي"""
    report = get_object_or_404(Report, uuid=uuid, created_by=request.user)
    
    return render(request, 'report_builder/builder.html', {
        'report': report,
    })


@login_required
@require_http_methods(['POST'])
def save_report(request, uuid):
    """حفظ التقرير"""
    report = get_object_or_404(Report, uuid=uuid, created_by=request.user)
    
    try:
        data = json.loads(request.body)
        
        report.name = data.get('name', report.name)
        report.description = data.get('description', report.description)
        report.report_type = data.get('report_type', report.report_type)
        report.selected_columns = data.get('columns', report.selected_columns)
        report.filters = data.get('filters', report.filters)
        report.sorting = data.get('sorting', report.sorting)
        report.grouping = data.get('grouping', report.grouping)
        report.calculations = data.get('calculations', report.calculations)
        report.config = data.get('config', report.config)
        report.styling = data.get('styling', report.styling)
        report.save()
        
        # حفظ العناصر
        if 'widgets' in data:
            report.widgets.all().delete()
            for i, w in enumerate(data['widgets']):
                ReportWidget.objects.create(
                    report=report,
                    widget_type=w.get('type'),
                    title=w.get('title', ''),
                    position_x=w.get('x', 0),
                    position_y=w.get('y', 0),
                    width=w.get('width', 6),
                    height=w.get('height', 4),
                    config=w.get('config', {}),
                    data_config=w.get('data_config', {}),
                    styling=w.get('styling', {}),
                    order=i
                )
        
        return JsonResponse({'success': True})
        
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})


@login_required
def preview_report(request, uuid):
    """معاينة التقرير"""
    report = get_object_or_404(Report, uuid=uuid)
    
    # التحقق من الصلاحية
    if report.created_by != request.user and request.user not in report.shared_with.all():
        if not report.is_public:
            messages.error(request, 'ليس لديك صلاحية')
            return redirect('report_builder:dashboard')
    
    # تحديث الإحصائيات
    report.view_count += 1
    report.last_viewed = timezone.now()
    report.save(update_fields=['view_count', 'last_viewed'])
    
    # تنفيذ التقرير
    try:
        service = ReportExecutionService(report)
        data = list(service.execute()[:100])  # تحديد أول 100 سجل للمعاينة
    except Exception as e:
        data = []
        messages.warning(request, f'خطأ في تنفيذ التقرير: {str(e)}')
    
    return render(request, 'report_builder/preview.html', {
        'report': report,
        'data': data,
    })


@login_required
@require_http_methods(['POST'])
def execute_report(request, uuid):
    """تنفيذ التقرير وإرجاع البيانات"""
    report = get_object_or_404(Report, uuid=uuid)
    
    try:
        parameters = json.loads(request.body) if request.body else {}
        service = ReportExecutionService(report)
        data = list(service.execute(parameters))
        
        return JsonResponse({
            'success': True,
            'data': data,
            'count': len(data)
        })
        
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})


@login_required
def export_report(request, uuid):
    """تصدير التقرير"""
    report = get_object_or_404(Report, uuid=uuid)
    format_type = request.GET.get('format', 'excel')
    
    service = ReportExecutionService(report)
    data = list(service.execute())
    
    if format_type == 'csv':
        import csv
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{report.name}.csv"'
        
        if data:
            writer = csv.DictWriter(response, fieldnames=data[0].keys())
            writer.writeheader()
            writer.writerows(data)
        
        return response
    
    elif format_type == 'excel':
        try:
            import openpyxl
            from io import BytesIO
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = report.name[:31]
            
            if data:
                # الرؤوس
                for col, header in enumerate(data[0].keys(), 1):
                    ws.cell(row=1, column=col, value=header)
                
                # البيانات
                for row_num, row_data in enumerate(data, 2):
                    for col, value in enumerate(row_data.values(), 1):
                        ws.cell(row=row_num, column=col, value=value)
            
            buffer = BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            
            response = HttpResponse(
                buffer.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = f'attachment; filename="{report.name}.xlsx"'
            return response
            
        except ImportError:
            messages.error(request, 'مكتبة openpyxl غير مثبتة')
            return redirect('report_builder:preview_report', uuid=uuid)
    
    return redirect('report_builder:preview_report', uuid=uuid)


@login_required
@require_http_methods(['DELETE'])
def delete_report(request, uuid):
    """حذف التقرير"""
    report = get_object_or_404(Report, uuid=uuid, created_by=request.user)
    report.delete()
    return JsonResponse({'success': True})


@login_required
def get_model_fields(request, content_type_id):
    """الحصول على حقول النموذج"""
    from django.contrib.contenttypes.models import ContentType
    
    try:
        ct = ContentType.objects.get(id=content_type_id)
        model = ct.model_class()
        fields = DataSourceService.get_model_fields(model)
        
        return JsonResponse({
            'success': True,
            'fields': fields
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})
